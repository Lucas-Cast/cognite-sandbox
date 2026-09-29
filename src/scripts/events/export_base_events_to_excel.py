"""Export BaseEvents to Excel with complete UTC timestamps.

The filter uses the event ``startTime``. Its default range includes every
event that starts from 10 September 2026 through 18 September 2026 UTC.
Edit the constants below before running:

    uv run python src/scripts/events/export_base_events_to_excel.py
"""

from __future__ import annotations

import json
import sys
from collections.abc import Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from cognite.client import CogniteClient
from cognite.client.data_classes.data_modeling import ViewId, filters
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# The script lives in src/scripts/events; core is a sibling package in src.
sys.path.insert(0, str(Path(__file__).parents[2]))

from core.cognite import create_cognite_client

# --- Configuration ---------------------------------------------------------
BASE_EVENT_VIEW = ViewId("sp_evt_glb_dmd", "BaseEvent", "v1.0.0")

# Leave None to export BaseEvents from every accessible instance space.
EVENT_INSTANCE_SPACE: str | None = "sp_otm_san_dat"

# The range is [START_TIME_GTE, END_TIME_LT): 10 through 18 September,
# inclusive, is represented by an exclusive bound at 19 September 00:00 UTC.
START_TIME_PROPERTY = "startTime"
START_TIME_GTE = "2026-09-10T00:00:00.000Z"
END_TIME_LT = "2026-09-19T00:00:00.000Z"

OUTPUT_FILENAME = "base_events_2026-09-10_to_2026-09-18_utc.xlsx"
WORKSHEET_NAME = "BaseEvents"
# ---------------------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parents[3]
OUTPUT_DIRECTORY = PROJECT_ROOT / "output"


def main() -> None:
    """Retrieve matching BaseEvents and export all view properties to Excel."""
    _validate_configuration()
    cognite_client = create_cognite_client()
    view_properties = _view_properties(cognite_client)
    if START_TIME_PROPERTY not in view_properties:
        available = ", ".join(sorted(view_properties))
        raise ValueError(
            f'Property "{START_TIME_PROPERTY}" is not available in '
            f"{_view_label(BASE_EVENT_VIEW)}. Available properties: {available}"
        )

    events = _retrieve_events(cognite_client)
    property_names = sorted(view_properties)
    output_path = _write_excel(events, property_names)

    print(f"BaseEvent view: {_view_label(BASE_EVENT_VIEW)}")
    print(f"Instance space: {EVENT_INSTANCE_SPACE or 'all accessible spaces'}")
    print(f"Filter: {START_TIME_PROPERTY} >= {START_TIME_GTE}")
    print(f"Filter: {START_TIME_PROPERTY} < {END_TIME_LT}")
    print(f"Events exported: {len(events)}")
    print(f"Excel saved to: {output_path}")


def _view_properties(cognite_client: CogniteClient) -> Mapping[str, Any]:
    views = cognite_client.data_modeling.views.retrieve(
        BASE_EVENT_VIEW,
        include_inherited_properties=True,
        all_versions=False,
    )
    if not views:
        raise ValueError(f"BaseEvent view not found: {_view_label(BASE_EVENT_VIEW)}")
    return views[0].properties


def _retrieve_events(cognite_client: CogniteClient) -> list[Any]:
    filters_to_apply: list[Any] = [
        filters.Range(
            _property_reference(START_TIME_PROPERTY),
            gte=START_TIME_GTE,
            lt=END_TIME_LT,
        ),
    ]
    if EVENT_INSTANCE_SPACE is not None:
        filters_to_apply.append(filters.Equals(["node", "space"], EVENT_INSTANCE_SPACE))

    return list(
        cognite_client.data_modeling.instances.list(
            instance_type="node",
            sources=BASE_EVENT_VIEW,
            filter=filters.And(*filters_to_apply),
            limit=-1,
        )
    )


def _write_excel(events: Sequence[Any], property_names: Sequence[str]) -> Path:
    OUTPUT_DIRECTORY.mkdir(exist_ok=True)
    output_path = OUTPUT_DIRECTORY / OUTPUT_FILENAME

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = WORKSHEET_NAME

    headers = [
        "node_space",
        "node_external_id",
        "node_version",
        "node_created_utc",
        "node_last_updated_utc",
        "node_type",
        *property_names,
    ]
    worksheet.append(headers)

    for event in events:
        properties = (
            event.properties.get(BASE_EVENT_VIEW, {}) if event.properties else {}
        )
        worksheet.append(
            [
                event.space,
                event.external_id,
                event.version,
                _format_epoch_milliseconds(getattr(event, "created_time", None)),
                _format_epoch_milliseconds(getattr(event, "last_updated_time", None)),
                _format_value(getattr(event, "type", None)),
                *[
                    _format_value(
                        properties.get(property_name),
                        is_timestamp=_is_timestamp_property(property_name),
                    )
                    for property_name in property_names
                ],
            ]
        )

    _format_worksheet(worksheet, headers)
    workbook.save(output_path)
    return output_path


def _format_worksheet(worksheet: Any, headers: Sequence[str]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.number_format = "@"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_index, header in enumerate(headers, start=1):
        width = max(14, min(60, len(header) + 2))
        if header.endswith("_utc") or _is_timestamp_property(header):
            width = 28
            for row_index in range(2, worksheet.max_row + 1):
                worksheet.cell(row_index, column_index).number_format = "@"
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def _property_reference(property_name: str) -> list[str]:
    return [
        BASE_EVENT_VIEW.space,
        f"{BASE_EVENT_VIEW.external_id}/{BASE_EVENT_VIEW.version}",
        property_name,
    ]


def _format_value(value: object, *, is_timestamp: bool = False) -> str | int | float:
    if value is None:
        return ""
    if is_timestamp:
        timestamp = _format_datetime(value)
        if timestamp is not None:
            return timestamp
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (str, int, float)):
        return value
    if isinstance(value, datetime):
        return _format_datetime(value) or ""
    return json.dumps(_json_value(value), ensure_ascii=False, separators=(",", ":"))


def _format_epoch_milliseconds(value: object) -> str:
    if not isinstance(value, int):
        return ""
    return (
        datetime.fromtimestamp(value / 1_000, tz=UTC).strftime(
            "%Y-%m-%d %H:%M:%S.%f UTC"
        )[:-7]
        + " UTC"
    )


def _format_datetime(value: object) -> str | None:
    if isinstance(value, datetime):
        timestamp = value if value.tzinfo else value.replace(tzinfo=UTC)
    elif isinstance(value, str):
        try:
            timestamp = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError:
            return None
        if timestamp.tzinfo is None:
            timestamp = timestamp.replace(tzinfo=UTC)
    else:
        return None
    return timestamp.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S.%f UTC")[:-3] + " UTC"


def _json_value(value: object) -> object:
    if isinstance(value, Mapping):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if isinstance(value, datetime):
        return _format_datetime(value)
    if hasattr(value, "space") and hasattr(value, "external_id"):
        return {
            "space": getattr(value, "space"),
            "externalId": getattr(value, "external_id"),
        }
    return value


def _is_timestamp_property(property_name: str) -> bool:
    normalized = "".join(
        character.lower() for character in property_name if character.isalnum()
    )
    return normalized.endswith(("time", "date", "at"))


def _view_label(view: ViewId) -> str:
    return f"{view.space}/{view.external_id}/{view.version}"


def _validate_configuration() -> None:
    if EVENT_INSTANCE_SPACE is not None and not EVENT_INSTANCE_SPACE.strip():
        raise ValueError("Set EVENT_INSTANCE_SPACE to a non-empty value or None.")
    if not START_TIME_PROPERTY.strip():
        raise ValueError("Set START_TIME_PROPERTY.")
    if _parse_utc(START_TIME_GTE) >= _parse_utc(END_TIME_LT):
        raise ValueError("START_TIME_GTE must be earlier than END_TIME_LT.")
    if Path(OUTPUT_FILENAME).name != OUTPUT_FILENAME or not OUTPUT_FILENAME.endswith(
        ".xlsx"
    ):
        raise ValueError("OUTPUT_FILENAME must be an .xlsx filename without a path.")


def _parse_utc(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(UTC)


if __name__ == "__main__":
    main()
